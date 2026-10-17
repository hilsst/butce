"""BUTCE_2023.xlsx içindeki Temmuz–Ekim 2026 sayfalarını düzenli JSON'a aktarır.

Kullanım: python3 aktar.py BUTCE_2023.xlsx butce.json
Aktarım sonunda Excel'in kendi hesapladığı toplamlarla karşılaştırma yapar.
"""
import json
import re
import sys

import openpyxl

AYLAR = {
    "TEMMUZ 2026": "2026-07",
    "AĞUSTOS 2026": "2026-08",
    "EYLÜL 2026": "2026-09",
    "EKİM 2026": "2026-10",
}
HAFTALAR = ["01-07", "08-14", "15-21", "22-28", "29-31"]
HAFTA_SUTUN = "JKLMN"

# Sabit gider adı -> kategori
KATEGORI = {
    "AXESS KART": "kredi_karti", "WORLDCARD": "kredi_karti", "BONUSKART": "kredi_karti",
    "AİDAT": "ev", "ARABA Bakım": "araba", "YAKIT": "araba",
    "ELEKTRİK": "fatura", "SU": "fatura", "İNTERNET+ KABLO TV": "fatura",
    "CEP TEL": "fatura", "İLKİM+İDİL CEP TEL": "fatura", "DİDEM CEP TEL": "fatura",
    "Yatırım": "birikim", "Bordo İkon": "birikim",
    "EV KREDİSİ": "kredi", "ARABA KREDİSİ": "kredi",
}


def sayi(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def gun(v):
    m = re.search(r"\d+", str(v)) if isinstance(v, str) else None
    return int(m.group()) if m else None


def ay_oku(ws):
    gelirler = []
    for r in range(3, 6):
        ad, tutar = ws[f"A{r}"].value, ws[f"B{r}"].value
        if ad:
            gelirler.append({"ad": ad.strip(), "tutar": sayi(tutar)})
    cuzdan = sayi(ws["B6"].value)

    sabit = []
    for r in range(3, 18):
        ad = ws[f"D{r}"].value
        if not ad:
            continue
        ad = ad.strip()
        planlanan, gerceklesen = ws[f"F{r}"].value, ws[f"E{r}"].value
        if planlanan is None and gerceklesen is None:
            continue
        sabit.append({
            "ad": ad,
            "kategori": KATEGORI.get(ad, "diger"),
            "planlanan": sayi(planlanan),
            "gerceklesen": sayi(gerceklesen),
            "odeme_gunu": gun(ws[f"G{r}"].value),
        })

    def haftalik(harcanan_satir, plan_satir):
        return [
            {"hafta": h, "planlanan": sayi(ws[f"{c}{plan_satir}"].value),
             "harcanan": sayi(ws[f"{c}{harcanan_satir}"].value)}
            for h, c in zip(HAFTALAR, HAFTA_SUTUN)
        ]

    # Satır 21 "GENEL" haftalık genel masrafın toplamıdır; ayrı kalem olarak tutulmaz.
    kisisel = []
    for r in range(22, 33):
        ad = ws[f"D{r}"].value
        if ad and (ws[f"E{r}"].value or ws[f"F{r}"].value):
            kisisel.append({"ad": ad.strip(), "planlanan": sayi(ws[f"F{r}"].value),
                            "harcanan": sayi(ws[f"E{r}"].value)})

    banka = ws["B14"].value
    return {
        "gelirler": gelirler,
        "cuzdana_cekilen": cuzdan,
        "sabit_giderler": sabit,
        "haftalik_mutfak": haftalik(4, 5),
        "haftalik_genel": haftalik(9, 10),
        "kisisel": kisisel,
        "banka_bakiyesi": sayi(banka) if banka is not None else None,
    }


def ozet(ay):
    t = lambda xs, k: sum(x[k] for x in xs)
    gelir = t(ay["gelirler"], "tutar") - ay["cuzdana_cekilen"]
    plan = (t(ay["sabit_giderler"], "planlanan") + t(ay["kisisel"], "planlanan")
            + t(ay["haftalik_genel"], "planlanan") + t(ay["haftalik_mutfak"], "planlanan"))
    reel = (t(ay["sabit_giderler"], "gerceklesen") + t(ay["kisisel"], "harcanan")
            + t(ay["haftalik_genel"], "harcanan") + t(ay["haftalik_mutfak"], "harcanan"))
    return {"gelir": gelir, "plan_harcanan": plan, "plan_kalan": gelir - plan,
            "reel_harcanan": reel, "mevcut": gelir - reel}


def main(kaynak, hedef):
    wb = openpyxl.load_workbook(kaynak, data_only=True)
    aylar, hata = {}, 0
    for sayfa, anahtar in AYLAR.items():
        ws = wb[sayfa]
        ay = ay_oku(ws)
        aylar[anahtar] = ay
        o = ozet(ay)
        excel = {"gelir": ws["B7"].value, "plan_harcanan": ws["B10"].value,
                 "plan_kalan": ws["B11"].value, "reel_harcanan": ws["B12"].value,
                 "mevcut": ws["B13"].value}
        for k, v in excel.items():
            durum = "OK" if abs(o[k] - sayi(v)) < 0.01 else "FARKLI"
            hata += durum != "OK"
            print(f"{anahtar} {k:14} excel={sayi(v):>10,.0f} uygulama={o[k]:>10,.0f} {durum}")
    with open(hedef, "w", encoding="utf-8") as f:
        json.dump({"surum": 1, "para_birimi": "TRY", "aylar": aylar}, f,
                  ensure_ascii=False, indent=2)
    print("Uyuşmayan toplam:", hata)


if __name__ == "__main__":
    main(*sys.argv[1:3])
